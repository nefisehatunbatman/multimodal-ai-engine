from __future__ import annotations

import hashlib
import logging
import asyncio
from io import BytesIO
from pathlib import Path

import httpx
from fastapi import UploadFile

from app.core.config import settings

logger = logging.getLogger(__name__)

WEKNORA_HOST = getattr(settings, "WEKNORA_APP_HOST", "localhost")
WEKNORA_PORT = int(getattr(settings, "WEKNORA_APP_PORT", 8080))
WEKNORA_API_KEY = getattr(settings, "WEKNORA_API_KEY", None)
WEKNORA_KB_ID = getattr(settings, "WEKNORA_KB_ID", None)

WEKNORA_BASE_URL = f"http://{WEKNORA_HOST}:{WEKNORA_PORT}"

ALLOWED_MIME_TYPES = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

ALLOWED_IMAGE_MIME_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

DOCX_MANUAL_CHUNK_SIZE = 700
DOCX_MANUAL_CHUNK_OVERLAP = 80


def _weknora_headers() -> dict[str, str]:
    return {"X-API-Key": WEKNORA_API_KEY or ""}


def _compute_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _extract_docx_text(file_bytes: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(file_bytes))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n\n".join(parts).strip()


def _split_text_for_manual_knowledge(text: str) -> str:
    paragraphs = [p.strip() for p in text.splitlines() if p.strip()]
    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        if len(paragraph) > DOCX_MANUAL_CHUNK_SIZE:
            if current:
                chunks.append(current.strip())
                current = ""
            start = 0
            while start < len(paragraph):
                end = start + DOCX_MANUAL_CHUNK_SIZE
                chunks.append(paragraph[start:end].strip())
                start = end
            continue

        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= DOCX_MANUAL_CHUNK_SIZE:
            current = candidate
        else:
            if current:
                chunks.append(current.strip())
            current = paragraph

    if current:
        chunks.append(current.strip())

    return "\n\n---\n\n".join(chunk for chunk in chunks if chunk)


async def _get_knowledge(client: httpx.AsyncClient, knowledge_id: str) -> dict:
    resp = await client.get(
        f"{WEKNORA_BASE_URL}/api/v1/knowledge/{knowledge_id}",
        headers=_weknora_headers(),
    )
    resp.raise_for_status()
    data = resp.json()
    return data.get("data") or {}


async def _enable_knowledge(client: httpx.AsyncClient, knowledge_id: str) -> None:
    resp = await client.put(
        f"{WEKNORA_BASE_URL}/api/v1/knowledge/{knowledge_id}",
        headers={**_weknora_headers(), "Content-Type": "application/json"},
        json={"enable_status": "enabled"},
    )
    resp.raise_for_status()


async def _reparse_knowledge(client: httpx.AsyncClient, knowledge_id: str) -> None:
    resp = await client.post(
        f"{WEKNORA_BASE_URL}/api/v1/knowledge/{knowledge_id}/reparse",
        headers=_weknora_headers(),
    )
    resp.raise_for_status()


async def wait_for_knowledge_ready(
    knowledge_id: str,
    *,
    timeout_seconds: int = 120,
    poll_seconds: float = 2.0,
) -> dict:
    """Wait until WeKnora finishes parsing/indexing a knowledge item."""
    deadline = asyncio.get_running_loop().time() + timeout_seconds
    last_item: dict = {}

    async with httpx.AsyncClient(timeout=30) as client:
        while True:
            last_item = await _get_knowledge(client, knowledge_id)
            parse_status = str(last_item.get("parse_status") or "")
            enable_status = str(last_item.get("enable_status") or "")
            error_message = str(last_item.get("error_message") or "")

            if parse_status == "completed":
                if enable_status != "enabled":
                    await _enable_knowledge(client, knowledge_id)
                    last_item = await _get_knowledge(client, knowledge_id)
                return last_item

            if parse_status == "failed":
                raise RuntimeError(
                    "WeKnora belge indeksleme basarisiz: "
                    f"{error_message or 'bilinmeyen hata'}"
                )

            if asyncio.get_running_loop().time() >= deadline:
                raise RuntimeError(
                    "WeKnora belge indeksleme zaman asimina ugradi "
                    f"(son durum: {parse_status or 'unknown'})"
                )

            await asyncio.sleep(poll_seconds)


async def ingest_file_to_weknora(file: UploadFile, user_id: int) -> dict:
    """
    Dosyayi WeKnora KB'ye yukler ve knowledge_id dondurur.
    user_id / knowledge_id eslesmesi cagiran tarafindan DB'ye kaydedilir.
    """
    if not WEKNORA_API_KEY:
        raise RuntimeError("WEKNORA_API_KEY is missing in environment")
    if not WEKNORA_KB_ID:
        raise RuntimeError("WEKNORA_KB_ID is missing in environment")

    suffix = Path(file.filename or "").suffix.lower()
    mime_type = ALLOWED_MIME_TYPES.get(suffix, "application/octet-stream")

    file_bytes = await file.read()
    if not file_bytes:
        raise RuntimeError("Yuklenen dosya bos")

    content_hash = _compute_hash(file_bytes)

    if suffix == ".docx":
        extracted_text = _extract_docx_text(file_bytes)
        if not extracted_text:
            raise RuntimeError("DOCX icinden okunabilir metin cikarilamadi")
        manual_content = (
            f"[DOCX Dosyasi: {file.filename}]\n\n"
            f"{_split_text_for_manual_knowledge(extracted_text)}"
        )
        upload_url = f"{WEKNORA_BASE_URL}/api/v1/knowledge-bases/{WEKNORA_KB_ID}/knowledge/manual"
        request_kwargs = {
            "headers": {**_weknora_headers(), "Content-Type": "application/json"},
            "json": {"title": file.filename or "DOCX belge", "content": manual_content},
        }
    else:
        upload_url = f"{WEKNORA_BASE_URL}/api/v1/knowledge-bases/{WEKNORA_KB_ID}/knowledge/file"
        request_kwargs = {
            "headers": _weknora_headers(),
            "files": {"file": (file.filename, file_bytes, mime_type)},
        }

    try:
        async with httpx.AsyncClient(timeout=120) as client:
            resp = await client.post(upload_url, **request_kwargs)
            resp.raise_for_status()
            data = resp.json()
    except httpx.HTTPStatusError as e:
        detail = e.response.text if e.response is not None else str(e)
        raise RuntimeError(f"WeKnora upload error (status={e.response.status_code}): {detail}")
    except httpx.HTTPError as e:
        raise RuntimeError(f"WeKnora unreachable: {e}")

    knowledge_id = str(data.get("data", {}).get("id", ""))
    if not knowledge_id:
        raise RuntimeError("WeKnora knowledge_id dondurmedi")

    if suffix == ".docx":
        async with httpx.AsyncClient(timeout=30) as client:
            await _reparse_knowledge(client, knowledge_id)

    ready_item = await wait_for_knowledge_ready(knowledge_id)

    logger.info("WeKnora ingestion complete: file=%s knowledge_id=%s hash=%s",
                file.filename, knowledge_id, content_hash[:16])

    return {
        "ok": True,
        "filename": file.filename,
        "knowledge_id": knowledge_id,
        "kb_id": WEKNORA_KB_ID,
        "source": "weknora",
        "parse_status": ready_item.get("parse_status"),
        "enable_status": ready_item.get("enable_status"),
        "content_hash": content_hash,
        "weknora_response": data,
    }


async def _generate_image_description(image_bytes: bytes, mime_type: str, filename: str) -> str:
    import base64
    from app.services.vision import call_openrouter_vision

    image_b64 = base64.b64encode(image_bytes).decode("utf-8")
    messages = [
        {"role": "system", "content": (
            "Sen bir gorsel analiz uzmanisın. Gorseldeki her detayi kapsamli sekilde acikla. "
            "Nesneleri, renkleri, konumlari, metinleri ve baglami detayliyla belirt. "
            "Aciklamanı Turkce yaz."
        )},
        {"role": "user", "content": [
            {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{image_b64}"}},
            {"type": "text", "text": "Bu gorseli tum detaylariyla acikla. Gorsel indeksleme icin kullanilacak."},
        ]},
    ]
    data = await call_openrouter_vision(messages)
    choices = data.get("choices", [])
    if not choices:
        raise RuntimeError("Vision model yanit vermedi")
    return choices[0].get("message", {}).get("content", "").strip()


async def ingest_image_to_weknora(image_bytes: bytes, mime_type: str, filename: str, user_id: int) -> dict:
    """
    Gorseli LLM ile analiz eder, aciklamayi WeKnora'ya text olarak kaydeder.
    knowledge_id dondurur — cagiran tarafindan DB'ye kaydedilir.
    """
    if not WEKNORA_API_KEY:
        raise RuntimeError("WEKNORA_API_KEY is missing in environment")
    if not WEKNORA_KB_ID:
        raise RuntimeError("WEKNORA_KB_ID is missing in environment")

    content_hash = _compute_hash(image_bytes)

    logger.info("Generating image description for: %s", filename)
    description = await _generate_image_description(image_bytes, mime_type, filename)
    if not description:
        raise RuntimeError("Gorsel aciklamasi uretilemedi")

    text_content = f"[Gorsel Dosyasi: {filename}]\n\n[Gorsel Aciklamasi]\n{description}"
    upload_url = f"{WEKNORA_BASE_URL}/api/v1/knowledge-bases/{WEKNORA_KB_ID}/knowledge/manual"

    try:
        async with httpx.AsyncClient(timeout=120) as client:
            resp = await client.post(
                upload_url,
                headers={**_weknora_headers(), "Content-Type": "application/json"},
                json={"title": f"Gorsel: {filename}", "content": text_content},
            )
            resp.raise_for_status()
            data = resp.json()
    except httpx.HTTPStatusError as e:
        detail = e.response.text if e.response is not None else str(e)
        raise RuntimeError(f"WeKnora image ingestion error (status={e.response.status_code}): {detail}")
    except httpx.HTTPError as e:
        raise RuntimeError(f"WeKnora unreachable: {e}")

    knowledge_id = str(data.get("data", {}).get("id", ""))
    if not knowledge_id:
        raise RuntimeError("WeKnora knowledge_id dondirmedi")

    try:
        async with httpx.AsyncClient(timeout=30) as client:
            await _reparse_knowledge(client, knowledge_id)
        ready_item = await wait_for_knowledge_ready(knowledge_id)
    except Exception as e:
        logger.warning("WeKnora reparse/enable failed: %s", e)
        ready_item = {}

    logger.info("Image ingested to WeKnora: file=%s knowledge_id=%s", filename, knowledge_id)

    return {
        "ok": True,
        "filename": filename,
        "knowledge_id": knowledge_id,
        "kb_id": WEKNORA_KB_ID,
        "source": "weknora_vision",
        "parse_status": ready_item.get("parse_status"),
        "enable_status": ready_item.get("enable_status"),
        "description_length": len(description),
        "content_hash": content_hash,
        "weknora_response": data,
    }


class DuplicateFileError(Exception):
    pass
